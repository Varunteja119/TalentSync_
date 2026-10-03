#!/usr/bin/env python
"""
Phase 3, step 5: human-labeled held-out check on the Vanetik & Kogan (2023) dataset
("Job Vacancy Ranking with Sentence Embeddings, Keywords, and Named Entities", Information 14(8):468).

    py -3.12 -m pip install python-docx
    py -3.12 ml\\natalia_eval.py
    (expects the clone at ..\\vacancy-resume-matching-dataset and the model at models\\minilm-resume-jd-ft)

PROTOCOL (fixed before any model was run on this data)
  * Annotations are RANK ARRAYS: index i = vacancy i (CSV row order), value = rank (1 = best).
    The paper averages the two annotators' arrays per resume, which only makes sense for this encoding.
  * Primary analysis: the 28 CVs where BOTH annotators gave valid permutations (annotator 1's rows for
    CV 9 and CV 28 contain duplicate ranks and are excluded). Truth = mean of the two rank arrays.
  * Primary input view: 'raw' (full resume text, full vacancy description; the model truncates at 256 tokens).
    Secondary view 'title': vacancy = "Role: <job_title>." only.
  * Sensitivity: all 30 CVs with raw arrays averaged as the paper does; each annotator alone; and
    human-vs-human (one annotator ranking against the other) as a reference for label noise.
  * Metrics: same harness as the CV experiments (pairwise accuracy, Spearman, top-1), bootstrap over CVs.
  * Caveats stated up front: n is tiny, annotators agree near chance, domain is IT-only with 5 similar
    vacancies, and the fine-tuned model was trained on compact views of GPT-parsed fields that cannot be
    rebuilt from free-text docx files. A null result here is NOT evidence against the model.

Outputs ml/results/natalia_eval.md / .json (metrics only; no resume text is printed or saved).
Data license: see the LICENSE file in the dataset repo. Keep the clone outside this repo and cite the paper.
"""

import argparse
import csv
import json
import os
import re
import sys

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from baseline_cv import table  # noqa: E402
from eval import (center_rel_by_jd, evaluate, evaluate_centered, paired_bootstrap,  # noqa: E402
                  random_baseline, spearman, summarize)
from inspect_netsol import Tee, hr  # noqa: E402

COLS = ["pairwise", "spearman", "top1"]
N_VAC = 5


# --------------------------------------------------------------------------- #
# Loading
# --------------------------------------------------------------------------- #
def parse_annotations(path):
    txt = open(path, "r", encoding="utf-8", errors="replace").read()
    i1, i2 = txt.index("ANNOTATOR_1_RANKINGS"), txt.index("ANNOTATOR_2_RANKINGS")
    pat = re.compile(r"\[\s*(\d)\s*,\s*(\d)\s*,\s*(\d)\s*,\s*(\d)\s*,\s*(\d)\s*\]")
    a1 = [[int(x) for x in m.groups()] for m in pat.finditer(txt[i1:i2])]
    a2 = [[int(x) for x in m.groups()] for m in pat.finditer(txt[i2:])]
    assert len(a1) == 30 and len(a2) == 30, f"expected 30 arrays each, got {len(a1)} and {len(a2)}"
    return a1, a2


def read_cv(path):
    if path.lower().endswith(".txt"):
        return open(path, "r", encoding="utf-8", errors="replace").read()
    try:
        import docx
    except ImportError:
        print("ERROR: pip install python-docx")
        sys.exit(1)
    d = docx.Document(path)
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return re.sub(r"[ \t]+", " ", "\n".join(parts)).strip()


def load_vacancies(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        rows = list(csv.DictReader(f))
    assert len(rows) == N_VAC, f"expected {N_VAC} vacancies, found {len(rows)}"
    return rows


def is_perm(a):
    return sorted(a) == list(range(1, N_VAC + 1))


# --------------------------------------------------------------------------- #
# Evaluation helpers
# --------------------------------------------------------------------------- #
def build_pairs(cvs, ranks):
    """ranks: {cv: [rank of vacancy 1..5]} -> pair dicts the shared harness understands."""
    return [{"id": f"{cv}_{v + 1}", "rh": f"cv{cv}", "jh": f"vac{v + 1}",
             "rel": (N_VAC - ranks[cv][v]) / (N_VAC - 1)} for cv in cvs for v in range(N_VAC)]


def aligned(cvs, score):
    return [score[(cv, v)] for cv in cvs for v in range(N_VAC)]


def human_scores(cvs, other_ranks):
    return {(cv, v): -float(other_ranks[cv][v]) for cv in cvs for v in range(N_VAC)}


def run_setting(title, cvs, truth, model_scores, n_boot, extra=None, paired=("fine-tuned", "pretrained")):
    pairs = build_pairs(cvs, truth)
    per = {name: evaluate(pairs, aligned(cvs, sc), "resume") for name, sc in model_scores.items()}
    per["random (chance)"] = random_baseline(pairs, "resume")
    for name, sc in (extra or {}).items():
        per[name] = evaluate(pairs, aligned(cvs, sc), "resume")
    rows = [(k, summarize(v, n_boot)) for k, v in per.items()]
    t = table(rows, COLS)
    hr(f"{title}   (n = {len(cvs)} CVs)")
    print("  " + t.replace("\n", "\n  "))
    md = f"\n## {title} (n = {len(cvs)} CVs)\n\n{t}"
    res = {"n": len(cvs), "rows": {k: s for k, s in rows}, "paired": []}
    a, b = paired
    if a in per and b in per:
        print(f"  paired: {a} minus {b}  (* = CI excludes 0)")
        md += f"\nPaired {a} minus {b}:\n\n| metric | diff [95% CI] |\n|---|---|\n"
        for m in COLS:
            d = paired_bootstrap(per[a], per[b], m)
            if d:
                star = "*" if d["significant"] else " "
                print(f"  {star} {m:9s} {d['mean_diff']:+.3f} [{d['lo']:+.3f}, {d['hi']:+.3f}]  n={d['n']}")
                md += f"| {m} | {d['mean_diff']:+.3f} [{d['lo']:+.3f}, {d['hi']:+.3f}]{star} |\n"
                res["paired"].append({"metric": m, **d})
    return md, res


# --------------------------------------------------------------------------- #
def diagnostics(valid, consensus, r1, r2, vacs, scores, tfidf_scores, n_boot):
    """EXPLORATORY / POST HOC (added after seeing the primary result). Separates the vacancy-level effect
    (which of the 5 vacancies tends to rank high) from resume-specific fit."""
    md, res = "\n## EXPLORATORY diagnostics (post hoc, not part of the pre-specified protocol)\n", {}
    hr("EXPLORATORY: vacancy-level view (post hoc)")
    names = list(scores["raw"].keys())
    mean_rank = lambda r: [float(np.mean([r[c][v] for c in valid])) for v in range(N_VAC)]
    mr1, mr2, mrc = mean_rank(r1), mean_rank(r2), mean_rank(consensus)
    mean_cos = {n: [float(np.mean([scores["raw"][n][(c, v)] for c in valid])) for v in range(N_VAC)] for n in names}
    head = "| vacancy | title | mean rank A1 | mean rank A2 | mean rank consensus | " + " | ".join(f"mean cosine {n}" for n in names) + " |\n"
    head += "|---|---|---|---|---|" + "---|" * len(names) + "\n"
    rows = ""
    for v in range(N_VAC):
        rows += (f"| {v + 1} | {vacs[v]['job_title'][:40]} | {mr1[v]:.2f} | {mr2[v]:.2f} | {mrc[v]:.2f} | "
                 + " | ".join(f"{mean_cos[n][v]:.3f}" for n in names) + " |\n")
    print("  " + (head + rows).replace("\n", "\n  "))
    md += "\n### Vacancy-level view\n\n" + head + rows
    res["vacancy_level"] = {"mean_rank_a1": mr1, "mean_rank_a2": mr2, "mean_rank_consensus": mrc,
                            "mean_cosine": mean_cos}
    for n in names:
        rho = spearman([-x for x in mrc], mean_cos[n])
        print(f"  vacancy-level Spearman, {n} mean cosine vs human (n=5 vacancies, descriptive only): "
              f"{rho if rho is None else round(rho, 3)}")
        res.setdefault("vacancy_level_spearman", {})[n] = rho

    # vacancy-centered, within-CV ordering
    pairs = build_pairs(valid, consensus)

    def centered(sc):
        m = {v: np.mean([sc[(c, v)] for c in valid]) for v in range(N_VAC)}
        return {(c, v): sc[(c, v)] - m[v] for c in valid for v in range(N_VAC)}

    per = {n: evaluate_centered(pairs, aligned(valid, centered(sc)), "resume") for n, sc in scores["raw"].items()}
    per["TF-IDF cosine (reference)"] = evaluate_centered(pairs, aligned(valid, centered(tfidf_scores)), "resume")
    rc = random_baseline(center_rel_by_jd(pairs), "resume")
    rc.pop("ndcg", None)
    per["random (chance)"] = rc
    rows_ = [(k, summarize(v, n_boot)) for k, v in per.items()]
    t = table(rows_, COLS)
    hr("EXPLORATORY: vacancy-centered, within-CV ordering (resume-specific fit; post hoc)")
    print("  " + t.replace("\n", "\n  "))
    md += "\n### Vacancy-centered within-CV ordering\n\n" + t
    res["vacancy_centered"] = {k: s_ for k, s_ in rows_}
    if "fine-tuned" in per and "pretrained" in per:
        print("  paired: fine-tuned minus pretrained  (* = CI excludes 0)")
        for m in COLS:
            d = paired_bootstrap(per["fine-tuned"], per["pretrained"], m)
            if d:
                print(f"  {'*' if d['significant'] else ' '} {m:9s} {d['mean_diff']:+.3f} "
                      f"[{d['lo']:+.3f}, {d['hi']:+.3f}]  n={d['n']}")

    # per-vacancy ranking of CVs
    hr("EXPLORATORY: per-vacancy Spearman across the 28 CVs (does the model rank CVs like the humans, within one vacancy?)")
    pv = {}
    line_head = "  vacancy:        " + "  ".join(f"{v + 1:>6d}" for v in range(N_VAC))
    print(line_head)
    for n, sc in list(scores["raw"].items()) + [("TF-IDF", tfidf_scores)]:
        vals = []
        for v in range(N_VAC):
            rho = spearman([-consensus[c][v] for c in valid], [sc[(c, v)] for c in valid])
            vals.append(rho)
        pv[n] = vals
        print(f"  {n:14s}  " + "  ".join(f"{x:6.2f}" if x is not None else "   n/a" for x in vals))
    res["per_vacancy_spearman"] = pv
    print("  (n=28 per cell; with human labels this noisy, individual cells within about +/-0.35 are indistinguishable from 0)")
    return md, res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--data-dir", default=os.path.join("..", "vacancy-resume-matching-dataset"))
    ap.add_argument("--finetuned", default=os.path.join("models", "minilm-resume-jd-ft"))
    ap.add_argument("--base", default="all-MiniLM-L6-v2")
    ap.add_argument("--n-boot", type=int, default=1000)
    ap.add_argument("--out", default=os.path.join("ml", "results", "natalia_eval"))
    args = ap.parse_args()
    sys.stdout = Tee(args.out + ".txt")

    from sentence_transformers import SentenceTransformer
    from sklearn.feature_extraction.text import TfidfVectorizer

    a1, a2 = parse_annotations(os.path.join(args.data_dir, "annotations-for-the-first-30-vacancies.txt"))
    vacs = load_vacancies(os.path.join(args.data_dir, "5_vacancies.csv"))
    cv_dir = os.path.join(args.data_dir, "CV")
    cv_text = {}
    for cv in range(1, 31):
        for ext in (".docx", ".txt"):
            p = os.path.join(cv_dir, f"{cv}{ext}")
            if os.path.exists(p):
                cv_text[cv] = read_cv(p)
                break
    missing = [c for c in range(1, 31) if c not in cv_text]
    if missing:
        print(f"ERROR: missing CV files for {missing}")
        sys.exit(1)

    r1 = {cv: a1[cv - 1] for cv in range(1, 31)}
    r2 = {cv: a2[cv - 1] for cv in range(1, 31)}
    valid = [cv for cv in range(1, 31) if is_perm(r1[cv]) and is_perm(r2[cv])]
    bad = [cv for cv in range(1, 31) if cv not in valid]
    consensus = {cv: [(r1[cv][v] + r2[cv][v]) / 2 for v in range(N_VAC)] for cv in range(1, 31)}

    hr("DATA CHECKS")
    print(f"  vacancies={len(vacs)}  CVs loaded={len(cv_text)}  CVs with both annotations valid={len(valid)}  "
          f"excluded (malformed rows)={bad}")
    print(f"  vacancy titles (CSV order = vacancy 1..5): {[v['job_title'] for v in vacs]}")
    print(f"  CV text length (chars): median={int(np.median([len(t) for t in cv_text.values()]))}  "
          f"vacancy description length: {[len(v['job_description']) for v in vacs]}")
    rho = [spearman([-x for x in r1[c]], [-x for x in r2[c]]) for c in valid]
    rho = [x for x in rho if x is not None]
    print(f"  inter-annotator Spearman (28 CVs): mean={np.mean(rho):.3f}  (chance = 0 +/- ~{0.5 / np.sqrt(len(rho)):.2f})")
    print("  >> Labels are very noisy; the 'human vs human' rows below show the realistic ceiling.")

    # ---- scores: model x view --------------------------------------------------
    views = {"raw": (lambda cv: cv_text[cv], lambda v: vacs[v]["job_description"]),
             "title": (lambda cv: cv_text[cv], lambda v: f"Role: {vacs[v]['job_title']}.")}

    def cosine_scores(model, view):
        rf, vf = views[view]
        res_texts = [rf(cv) for cv in range(1, 31)]
        vac_texts = [vf(v) for v in range(N_VAC)]
        er = np.asarray(model.encode(res_texts, batch_size=16, normalize_embeddings=True, show_progress_bar=False))
        ev = np.asarray(model.encode(vac_texts, batch_size=16, normalize_embeddings=True, show_progress_bar=False))
        S = er @ ev.T
        return {(cv, v): float(S[cv - 1, v]) for cv in range(1, 31) for v in range(N_VAC)}

    models = {"pretrained": SentenceTransformer(args.base)}
    if os.path.isdir(args.finetuned):
        models["fine-tuned"] = SentenceTransformer(args.finetuned)
    else:
        print(f"  [!] fine-tuned model not found at {args.finetuned}; only pretrained will be scored.")
    scores = {view: {name: cosine_scores(m, view) for name, m in models.items()} for view in views}

    # TF-IDF reference (unsupervised, fitted on the 65 CVs + 5 vacancies; no labels used)
    all_cv = sorted(int(os.path.splitext(f)[0]) for f in os.listdir(cv_dir) if f.lower().endswith((".docx", ".txt")))
    corpus = [read_cv(os.path.join(cv_dir, f"{c}{'.docx' if os.path.exists(os.path.join(cv_dir, f'{c}.docx')) else '.txt'}"))
              for c in all_cv] + [v["job_description"] for v in vacs]
    tfidf = TfidfVectorizer(ngram_range=(1, 2), sublinear_tf=True, stop_words="english").fit(corpus)
    Xc = tfidf.transform([cv_text[c] for c in range(1, 31)])
    Xv = tfidf.transform([v["job_description"] for v in vacs])
    St = (Xc @ Xv.T).toarray()
    tfidf_scores = {(cv, v): float(St[cv - 1, v]) for cv in range(1, 31) for v in range(N_VAC)}

    md = f"# Human-labeled held-out check (Vanetik & Kogan 2023), n={len(valid)} CVs\n"
    md += ("\nCaveats: tiny n; annotators agree near chance; IT-only resumes and 5 similar vacancies; fine-tuned model "
           "trained on compact views of GPT-parsed fields that cannot be rebuilt from free text.\n")
    results = {"excluded_cvs": bad, "inter_annotator_spearman": float(np.mean(rho))}

    m_, r_ = run_setting("PRIMARY: raw view, consensus of both annotators", valid, consensus,
                         scores["raw"], args.n_boot, extra={"TF-IDF cosine (reference)": tfidf_scores})
    md += m_
    results["primary_raw"] = r_

    m_, r_ = run_setting("Sensitivity: raw view, all 30 CVs, raw arrays averaged (paper's approach)",
                         list(range(1, 31)), consensus, scores["raw"], args.n_boot)
    md += m_
    results["sens_all30"] = r_

    m_, r_ = run_setting("Sensitivity: 'title' view (vacancy = job title only), consensus", valid, consensus,
                         scores["title"], args.n_boot)
    md += m_
    results["sens_title"] = r_

    for name, truth, other in (("annotator 1", r1, r2), ("annotator 2", r2, r1)):
        m_, r_ = run_setting(f"Sensitivity: raw view vs {name} alone (human = the other annotator)",
                             valid, truth, scores["raw"], args.n_boot,
                             extra={"human vs human (other annotator)": human_scores(valid, other)})
        md += m_
        results[f"sens_{name.replace(' ', '')}"] = r_

    m_, r_ = diagnostics(valid, consensus, r1, r2, vacs, scores, tfidf_scores, args.n_boot)
    md += m_
    results["exploratory"] = r_

    with open(args.out + ".md", "w", encoding="utf-8") as f:
        f.write(md)
    with open(args.out + ".json", "w", encoding="utf-8") as f:
        json.dump(results, f, indent=2, default=float)
    print(f"\nSaved: {args.out}.md / .json")


if __name__ == "__main__":
    main()
